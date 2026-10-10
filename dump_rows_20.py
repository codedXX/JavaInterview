import openpyxl
import json

excel_path = r'C:\Users\YX\Desktop\2026年10月预付定金小程序抽奖.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb.active

with open(r'F:\Projects\JavaInterview\rows_20_25.txt', 'w', encoding='utf-8') as f:
    f.write(f"sheet.max_row: {sheet.max_row}\n")
    for r in range(20, sheet.max_row + 1):
        vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        f.write(f"R{r}: {vals}\n")
