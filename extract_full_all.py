import openpyxl
import json

wb = openpyxl.load_workbook(r'C:\Users\YX\Desktop\2026年10月预付定金小程序抽奖.xlsx', data_only=True)

result = {}
for name in wb.sheetnames:
    s = wb[name]
    s_data = []
    for r in range(1, s.max_row + 1):
        vals = [s.cell(r, c).value for c in range(1, s.max_column + 1)]
        if any(v is not None for v in vals):
            s_data.append({"row": r, "vals": [str(x) if x is not None else "" for x in vals]})
    result[name] = s_data

with open(r'F:\Projects\JavaInterview\preview_rows.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)
