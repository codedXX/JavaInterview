import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\YX\Desktop\2026年10月预付定金小程序抽奖.xlsx', data_only=True)
sheet = wb.active

non_empty = []
for r in range(20, sheet.max_row + 1):
    vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
    if any(x is not None for x in vals):
        non_empty.append((r, vals))

with open(r'F:\Projects\JavaInterview\preview_rows.json', 'w', encoding='utf-8') as f:
    import json
    json.dump(non_empty, f)
