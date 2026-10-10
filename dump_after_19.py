import openpyxl
import json

excel_path = r'C:\Users\YX\Desktop\2026年10月预付定金小程序抽奖.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb.active

res = []
for r in range(20, sheet.max_row + 1):
    vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
    res.append({"row": r, "vals": [str(x) if x is not None else "" for x in vals]})

with open(r'F:\Projects\JavaInterview\rows_after_19.json', 'w', encoding='utf-8') as f:
    json.dump({"max_row": sheet.max_row, "rows": res}, f, ensure_ascii=False, indent=2)
