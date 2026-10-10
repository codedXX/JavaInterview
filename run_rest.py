import openpyxl
import json

excel_path = r'C:\Users\YX\Desktop\2026年10月预付定金小程序抽奖.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb.active

data = []
for r in range(20, sheet.max_row + 1):
    row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
    data.append({"row": r, "vals": [str(x) if x is not None else "" for x in row_vals]})

out_path = r'F:\Projects\JavaInterview\preview_rows_rest.json'
with open(out_path, 'w', encoding='utf-8') as f:
    json.dump({"rows": data}, f, ensure_ascii=False, indent=2)
