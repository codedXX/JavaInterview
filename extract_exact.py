import openpyxl
import json

wb = openpyxl.load_workbook(r'C:\Users\YX\Desktop\2026年10月预付定金小程序抽奖.xlsx', data_only=True)
sheet = wb.active

products = []
for r in range(5, 20):
    pid = sheet.cell(r, 3).value
    name = sheet.cell(r, 4).value
    if pid is not None and name is not None:
        clean_pid = str(pid).strip()
        clean_name = " ".join(str(name).split())
        products.append({
            "productId": clean_pid,
            "productName": clean_name
        })

with open(r'F:\Projects\JavaInterview\preview_rows.json', 'w', encoding='utf-8') as f:
    json.dump(products, f, ensure_ascii=False, indent=2)
