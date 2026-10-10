import openpyxl
import json

excel_path = r'C:\Users\YX\Desktop\2026年10月预付定金小程序抽奖.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb.active

print(f"Max row: {sheet.max_row}")
for r in range(1, sheet.max_row + 1):
    vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
    print(f"R{r}: {vals}")
