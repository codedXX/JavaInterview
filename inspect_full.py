import openpyxl
import json

excel_path = r'C:\Users\YX\Desktop\2026年10月预付定金小程序抽奖.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

result = {}
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    sheet_data = []
    for r in range(1, sheet.max_row + 1):
        vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        if any(v is not None for v in vals):
            sheet_data.append({"row": r, "vals": [str(x) if x is not None else "" for x in vals]})
    result[sheet_name] = sheet_data

with open(r'F:\Projects\JavaInterview\all_sheets.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)
