import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\YX\Desktop\2026年10月预付定金小程序抽奖.xlsx', data_only=True)
sheet = wb.active
with open(r'F:\Projects\JavaInterview\excel_inspect.txt', 'w', encoding='utf-8') as f:
    f.write(f"Sheet title: {sheet.title}, max_row: {sheet.max_row}, max_column: {sheet.max_column}\n")
    for r in range(1, min(25, sheet.max_row + 1)):
        row_vals = [str(sheet.cell(r, c).value) for c in range(1, min(15, sheet.max_column + 1))]
        f.write(f"Row {r}: {row_vals}\n")
print("Done writing inspect")
